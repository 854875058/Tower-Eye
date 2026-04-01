from __future__ import annotations

import asyncio
import csv
import hashlib
import html
import json
import mimetypes
import os
import re
import shutil
import sqlite3
import stat
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path, PurePosixPath
from typing import Iterable, Optional
from xml.etree import ElementTree

import numpy as np
import pandas as pd
from fastapi import UploadFile
from sqlalchemy import case, delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.logging import get_logger
from app.models.models import (
    ProcessingStatus,
    WorkbenchChunk,
    WorkbenchDataset,
    WorkbenchResource,
    WorkbenchSourceKind,
    WorkbenchSubscription,
    WorkbenchSubscriptionStatus,
)
from app.schemas.schemas import WorkbenchSearchResult
from app.services.media_models import media_model_client

logger = get_logger(__name__)

BACKEND_DIR = Path(__file__).resolve().parents[2]
PROJECT_ROOT = Path(__file__).resolve().parents[3]
WORKBENCH_ROOT = (BACKEND_DIR / settings.WORKBENCH_STORAGE_DIR).resolve()
WORKBENCH_UPLOAD_ROOT = WORKBENCH_ROOT / "uploads"
WORKBENCH_PATH_ROOT = WORKBENCH_ROOT / "paths"
WORKBENCH_SFTP_ROOT = WORKBENCH_ROOT / "sftp"
WORKBENCH_EXTRACT_ROOT = WORKBENCH_ROOT / "extracted"
WORKBENCH_IMPORT_ROOT = WORKBENCH_ROOT / "imports"

TOWER_METADATA_DB = PROJECT_ROOT / "data" / "metadata.db"

ARCHIVE_EXTENSIONS = {".zip"}
TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".json",
    ".jsonl",
    ".yaml",
    ".yml",
    ".xml",
    ".html",
    ".htm",
    ".sql",
    ".log",
    ".csv",
    ".tsv",
}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xls"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx"}
SUPPORTED_EXTENSIONS = ARCHIVE_EXTENSIONS | TEXT_EXTENSIONS | SPREADSHEET_EXTENSIONS | DOCUMENT_EXTENSIONS

TEXT_STOP_WORDS = {
    "the",
    "and",
    "for",
    "with",
    "that",
    "this",
    "from",
    "have",
    "into",
    "your",
    "you",
    "are",
    "was",
    "were",
    "than",
    "then",
    "data",
    "file",
    "sheet",
    "row",
    "column",
    "name",
}

try:
    import paramiko
except Exception:  # pragma: no cover
    paramiko = None

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover
    DocxDocument = None

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover
    PdfReader = None


@dataclass
class PreparedWorkbenchFile:
    source_kind: WorkbenchSourceKind
    source_uri: str
    stored_path: str
    file_name: str
    file_extension: str
    file_category: str
    mime_type: Optional[str]
    file_size: int
    checksum: str
    extraction_metadata: dict = field(default_factory=dict)


@dataclass
class ExtractedChunks:
    parser_name: str
    file_category: str
    chunks: list[str]
    text_length: int
    metadata: dict = field(default_factory=dict)


def ensure_workbench_dirs() -> None:
    for path in (
        WORKBENCH_ROOT,
        WORKBENCH_UPLOAD_ROOT,
        WORKBENCH_PATH_ROOT,
        WORKBENCH_SFTP_ROOT,
        WORKBENCH_EXTRACT_ROOT,
        WORKBENCH_IMPORT_ROOT,
    ):
        path.mkdir(parents=True, exist_ok=True)


def _normalize_local_path(path: str) -> str:
    return str(Path(path).expanduser().resolve())


def _compute_checksum(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as file_obj:
        for chunk in iter(lambda: file_obj.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _guess_file_category(file_name: str) -> Optional[str]:
    suffix = Path(file_name).suffix.lower()
    if suffix in {".csv", ".tsv"}:
        return "table"
    if suffix in SPREADSHEET_EXTENSIONS:
        return "spreadsheet"
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".docx":
        return "docx"
    if suffix in {".json", ".jsonl"}:
        return "json"
    if suffix in {".html", ".htm"}:
        return "html"
    if suffix == ".xml":
        return "xml"
    if suffix in {".md", ".markdown"}:
        return "markdown"
    if suffix in ARCHIVE_EXTENSIONS:
        return "archive"
    if suffix in SUPPORTED_EXTENSIONS:
        return "text"
    return None


def _make_safe_name(file_name: str) -> str:
    safe = re.sub(r"[^0-9a-zA-Z._-]+", "_", Path(file_name).name)
    return safe or "file.bin"


def _copy_file_to_root(source_path: str, target_root: Path) -> PreparedWorkbenchFile:
    ensure_workbench_dirs()
    src = Path(source_path)
    if not src.exists() or not src.is_file():
        raise FileNotFoundError(source_path)

    target_root.mkdir(parents=True, exist_ok=True)
    target_path = target_root / f"{hashlib.md5(str(src).encode('utf-8')).hexdigest()[:8]}_{_make_safe_name(src.name)}"
    shutil.copy2(src, target_path)
    category = _guess_file_category(src.name)
    if not category:
        raise ValueError(f"Unsupported file type: {src.name}")
    return PreparedWorkbenchFile(
        source_kind=WorkbenchSourceKind.PATH,
        source_uri=str(src.resolve()),
        stored_path=str(target_path),
        file_name=src.name,
        file_extension=src.suffix.lower(),
        file_category=category,
        mime_type=mimetypes.guess_type(src.name)[0],
        file_size=int(target_path.stat().st_size),
        checksum=_compute_checksum(str(target_path)),
    )


async def _save_upload_file(upload: UploadFile, target_root: Path) -> PreparedWorkbenchFile:
    ensure_workbench_dirs()
    target_root.mkdir(parents=True, exist_ok=True)
    original_name = upload.filename or "upload.bin"
    category = _guess_file_category(original_name)
    if not category:
        raise ValueError(f"Unsupported file type: {original_name}")

    safe_name = f"{hashlib.md5(original_name.encode('utf-8')).hexdigest()[:8]}_{_make_safe_name(original_name)}"
    target_path = target_root / safe_name
    content = await upload.read()
    with open(target_path, "wb") as file_obj:
        file_obj.write(content)

    return PreparedWorkbenchFile(
        source_kind=WorkbenchSourceKind.UPLOAD,
        source_uri=original_name,
        stored_path=str(target_path),
        file_name=Path(original_name).name,
        file_extension=Path(original_name).suffix.lower(),
        file_category=category,
        mime_type=upload.content_type or mimetypes.guess_type(original_name)[0],
        file_size=len(content),
        checksum=hashlib.sha256(content).hexdigest(),
    )


def _iter_local_inputs(input_paths: Iterable[str]) -> list[str]:
    discovered: list[str] = []
    seen: set[str] = set()
    for raw_path in input_paths:
        if not raw_path:
            continue
        normalized = _normalize_local_path(raw_path)
        path_obj = Path(normalized)
        if not path_obj.exists():
            raise FileNotFoundError(normalized)

        candidates = sorted(path_obj.rglob("*")) if path_obj.is_dir() else [path_obj]
        for candidate in candidates:
            if not candidate.is_file():
                continue
            if candidate.suffix.lower() not in SUPPORTED_EXTENSIONS:
                continue
            resolved = str(candidate.resolve())
            if resolved in seen:
                continue
            seen.add(resolved)
            discovered.append(resolved)
    return discovered


def _is_safe_archive_member(name: str) -> bool:
    path = PurePosixPath(name)
    if path.is_absolute():
        return False
    return all(part not in {"..", ""} for part in path.parts)


def _extract_archive(file: PreparedWorkbenchFile, dataset_key: str) -> list[PreparedWorkbenchFile]:
    archive_path = Path(file.stored_path)
    target_root = WORKBENCH_EXTRACT_ROOT / dataset_key / archive_path.stem
    target_root.mkdir(parents=True, exist_ok=True)
    extracted: list[PreparedWorkbenchFile] = []

    with zipfile.ZipFile(archive_path, "r") as zip_file:
        for info in zip_file.infolist():
            if info.is_dir():
                continue
            if not _is_safe_archive_member(info.filename):
                continue
            member_name = Path(info.filename).name
            category = _guess_file_category(member_name)
            if not category or category == "archive":
                continue

            target_path = target_root / info.filename.replace("/", os.sep)
            target_path.parent.mkdir(parents=True, exist_ok=True)
            with zip_file.open(info, "r") as src, open(target_path, "wb") as dst:
                shutil.copyfileobj(src, dst)

            extracted.append(
                PreparedWorkbenchFile(
                    source_kind=file.source_kind,
                    source_uri=f"zip://{file.source_uri}!/{info.filename}",
                    stored_path=str(target_path.resolve()),
                    file_name=member_name,
                    file_extension=Path(member_name).suffix.lower(),
                    file_category=category,
                    mime_type=mimetypes.guess_type(member_name)[0],
                    file_size=int(target_path.stat().st_size),
                    checksum=_compute_checksum(str(target_path)),
                    extraction_metadata={
                        "archive_source": file.source_uri,
                        "archive_member": info.filename,
                    },
                )
            )
    return extracted


def _read_text_file(path: str) -> str:
    encodings = ["utf-8", "utf-8-sig", "gb18030", "latin-1"]
    for encoding in encodings:
        try:
            return Path(path).read_text(encoding=encoding)
        except Exception:
            continue
    with open(path, "rb") as file_obj:
        return file_obj.read().decode("utf-8", errors="ignore")


def _strip_markup(text: str) -> str:
    text = re.sub(r"<script[\s\S]*?</script>", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"<style[\s\S]*?</style>", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    return html.unescape(text)


def _chunk_text(text: str, *, chunk_size: int, overlap: int) -> list[str]:
    normalized = re.sub(r"\r\n?", "\n", text)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized).strip()
    if not normalized:
        return []

    chunks: list[str] = []
    start = 0
    length = len(normalized)
    while start < length:
        end = min(start + chunk_size, length)
        if end < length:
            split_at = normalized.rfind("\n", start, end)
            if split_at > start + chunk_size // 2:
                end = split_at
        chunk = normalized[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= length:
            break
        start = max(end - overlap, start + 1)
    return chunks


def _sanitize_cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    if len(text) > 160:
        text = text[:157] + "..."
    return text


def _build_tabular_chunks(
    rows: Iterable[tuple[int, dict[str, object]]],
    *,
    title: str,
    chunk_size: int,
) -> tuple[list[str], int]:
    chunks: list[str] = []
    current_lines: list[str] = [title]
    current_length = len(title)
    row_count = 0

    for row_no, row in rows:
        values = [f"{key}={_sanitize_cell(value)}" for key, value in row.items() if _sanitize_cell(value)]
        if not values:
            continue
        line = f"Row {row_no}: " + "; ".join(values)
        if current_length + len(line) + 1 > chunk_size and len(current_lines) > 1:
            chunks.append("\n".join(current_lines))
            current_lines = [title, line]
            current_length = len(title) + len(line) + 1
        else:
            current_lines.append(line)
            current_length += len(line) + 1
        row_count += 1

    if len(current_lines) > 1:
        chunks.append("\n".join(current_lines))
    return chunks, row_count


def _extract_csv_chunks(path: str, *, separator: str) -> ExtractedChunks:
    rows: list[tuple[int, dict[str, object]]] = []
    with open(path, "r", encoding="utf-8", errors="ignore", newline="") as file_obj:
        reader = csv.DictReader(file_obj, delimiter=separator)
        for row_no, row in enumerate(reader, start=1):
            rows.append((row_no, row))
    title = f"Table file: {Path(path).name}"
    chunks, row_count = _build_tabular_chunks(
        rows,
        title=title,
        chunk_size=max(int(settings.WORKBENCH_CHUNK_SIZE), 600),
    )
    return ExtractedChunks(
        parser_name="csv" if separator == "," else "tsv",
        file_category="table",
        chunks=chunks,
        text_length=sum(len(chunk) for chunk in chunks),
        metadata={
            "row_count": row_count,
            "column_count": len(rows[0][1].keys()) if rows else 0,
        },
    )


def _extract_excel_chunks(path: str) -> ExtractedChunks:
    excel_file = pd.ExcelFile(path)
    all_chunks: list[str] = []
    sheet_meta: list[dict] = []

    for sheet_name in excel_file.sheet_names:
        frame = excel_file.parse(sheet_name).fillna("")
        records = [(index + 1, row.to_dict()) for index, row in frame.iterrows()]
        title = f"Spreadsheet: {Path(path).name}\nSheet: {sheet_name}"
        chunks, row_count = _build_tabular_chunks(
            records,
            title=title,
            chunk_size=max(int(settings.WORKBENCH_CHUNK_SIZE), 600),
        )
        all_chunks.extend(chunks)
        sheet_meta.append(
            {
                "sheet_name": sheet_name,
                "row_count": row_count,
                "column_count": len(frame.columns),
            }
        )

    return ExtractedChunks(
        parser_name="openpyxl",
        file_category="spreadsheet",
        chunks=all_chunks,
        text_length=sum(len(chunk) for chunk in all_chunks),
        metadata={"sheet_count": len(sheet_meta), "sheets": sheet_meta},
    )


def _extract_docx_text(path: str) -> ExtractedChunks:
    if DocxDocument is None:
        raise RuntimeError("python-docx is not installed")

    document = DocxDocument(path)
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        table_lines = [f"Table {table_index}"]
        for row in table.rows:
            row_cells = [_sanitize_cell(cell.text) for cell in row.cells if _sanitize_cell(cell.text)]
            if row_cells:
                table_lines.append(" | ".join(row_cells))
        if len(table_lines) > 1:
            blocks.append("\n".join(table_lines))

    text = "\n\n".join(blocks)
    chunks = _chunk_text(
        text,
        chunk_size=int(settings.WORKBENCH_CHUNK_SIZE),
        overlap=int(settings.WORKBENCH_CHUNK_OVERLAP),
    )
    return ExtractedChunks(
        parser_name="python-docx",
        file_category="docx",
        chunks=chunks,
        text_length=len(text),
        metadata={
            "paragraph_count": len(document.paragraphs),
            "table_count": len(document.tables),
        },
    )


def _extract_pdf_text(path: str) -> ExtractedChunks:
    if PdfReader is None:
        raise RuntimeError("pypdf is not installed")

    reader = PdfReader(path)
    pages: list[str] = []
    for page_index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"Page {page_index}\n{text}")
    merged = "\n\n".join(pages)
    chunks = _chunk_text(
        merged,
        chunk_size=int(settings.WORKBENCH_CHUNK_SIZE),
        overlap=int(settings.WORKBENCH_CHUNK_OVERLAP),
    )
    return ExtractedChunks(
        parser_name="pypdf",
        file_category="pdf",
        chunks=chunks,
        text_length=len(merged),
        metadata={"page_count": len(reader.pages)},
    )


def _extract_json_text(path: str) -> ExtractedChunks:
    raw_text = _read_text_file(path).strip()
    if Path(path).suffix.lower() == ".json":
        try:
            raw_text = json.dumps(json.loads(raw_text), ensure_ascii=False, indent=2)
        except Exception:
            pass
    chunks = _chunk_text(
        raw_text,
        chunk_size=int(settings.WORKBENCH_CHUNK_SIZE),
        overlap=int(settings.WORKBENCH_CHUNK_OVERLAP),
    )
    return ExtractedChunks(
        parser_name="json",
        file_category="json",
        chunks=chunks,
        text_length=len(raw_text),
        metadata={"line_count": raw_text.count("\n") + 1 if raw_text else 0},
    )


def _extract_xml_text(path: str) -> ExtractedChunks:
    try:
        root = ElementTree.parse(path).getroot()
        text = " ".join(part.strip() for part in root.itertext() if part.strip())
    except Exception:
        text = _read_text_file(path)
    chunks = _chunk_text(
        text,
        chunk_size=int(settings.WORKBENCH_CHUNK_SIZE),
        overlap=int(settings.WORKBENCH_CHUNK_OVERLAP),
    )
    return ExtractedChunks(
        parser_name="xml",
        file_category="xml",
        chunks=chunks,
        text_length=len(text),
    )


def _extract_html_text(path: str) -> ExtractedChunks:
    text = _strip_markup(_read_text_file(path))
    chunks = _chunk_text(
        text,
        chunk_size=int(settings.WORKBENCH_CHUNK_SIZE),
        overlap=int(settings.WORKBENCH_CHUNK_OVERLAP),
    )
    return ExtractedChunks(
        parser_name="html",
        file_category="html",
        chunks=chunks,
        text_length=len(text),
    )


def _extract_plain_text(path: str, category: str) -> ExtractedChunks:
    text = _read_text_file(path)
    chunks = _chunk_text(
        text,
        chunk_size=int(settings.WORKBENCH_CHUNK_SIZE),
        overlap=int(settings.WORKBENCH_CHUNK_OVERLAP),
    )
    return ExtractedChunks(
        parser_name="text",
        file_category=category,
        chunks=chunks,
        text_length=len(text),
        metadata={"line_count": text.count("\n") + 1 if text else 0},
    )


def extract_resource_chunks(path: str) -> ExtractedChunks:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        return _extract_csv_chunks(path, separator=",")
    if suffix == ".tsv":
        return _extract_csv_chunks(path, separator="\t")
    if suffix in SPREADSHEET_EXTENSIONS:
        return _extract_excel_chunks(path)
    if suffix == ".docx":
        return _extract_docx_text(path)
    if suffix == ".pdf":
        return _extract_pdf_text(path)
    if suffix in {".json", ".jsonl"}:
        return _extract_json_text(path)
    if suffix == ".xml":
        return _extract_xml_text(path)
    if suffix in {".html", ".htm"}:
        return _extract_html_text(path)
    if suffix in {".md", ".markdown"}:
        return _extract_plain_text(path, "markdown")
    if suffix in TEXT_EXTENSIONS:
        return _extract_plain_text(path, "text")
    raise RuntimeError(f"Unsupported file type: {suffix or 'unknown'}")


def _extract_keywords(text: str, *, limit: int = 12) -> list[str]:
    tokens = re.findall(r"[\u4e00-\u9fff]+|[a-zA-Z0-9]+", text.lower())
    scores: dict[str, int] = {}
    for token in tokens:
        if len(token) <= 1 or token in TEXT_STOP_WORDS:
            continue
        scores[token] = scores.get(token, 0) + 1
    ranked = sorted(scores.items(), key=lambda item: (-item[1], item[0]))
    return [token for token, _count in ranked[:limit]]


def _cosine_similarity(left: list[float], right: list[float]) -> float:
    if not left or not right:
        return 0.0
    left_vec = np.asarray(left, dtype=np.float32)
    right_vec = np.asarray(right, dtype=np.float32)
    size = min(left_vec.size, right_vec.size)
    if size <= 0:
        return 0.0
    left_vec = left_vec[:size]
    right_vec = right_vec[:size]
    denom = float(np.linalg.norm(left_vec) * np.linalg.norm(right_vec))
    if denom <= 0:
        return 0.0
    return float(np.dot(left_vec, right_vec) / denom)


def _keyword_match_score(query: str, text: str) -> float:
    if not query or not text:
        return 0.0
    tokens = [token for token in re.findall(r"[\u4e00-\u9fff]+|[a-zA-Z0-9]+", query.lower()) if token]
    if not tokens:
        return 0.0
    corpus = text.lower()
    hit_count = sum(1 for token in tokens if token in corpus)
    return hit_count / max(len(tokens), 1)


def _sanitize_summary_text(value: Optional[str]) -> str:
    text = str(value or "").strip()
    text = text.replace("\r", " ").replace("\n", " ")
    text = re.sub(r"\s+", " ", text)
    return text


def _format_tower_event_markdown(row: sqlite3.Row) -> str:
    summary = _sanitize_summary_text(row["summary"])
    description = _sanitize_summary_text(row["description"])
    extra_json = row["extra_json"] or ""
    if extra_json and len(str(extra_json)) > 3000:
        extra_json = str(extra_json)[:3000] + "..."

    return "\n".join(
        [
            f"# 视联告警事件 {row['event_id']}",
            "",
            "## 基础信息",
            f"- 事件类型: {row['event_type'] or ''}",
            f"- 告警级别: {row['alarm_level'] or ''}",
            f"- 告警来源: {row['alarm_source'] or ''}",
            f"- 告警时间: {row['alarm_time'] or ''}",
            f"- 置信度: {row['confidence_level'] or ''}",
            f"- 算法名称: {row['algorithm_name'] or ''}",
            f"- 场景: {row['app_scenarios'] or ''}",
            "",
            "## 位置与设备",
            f"- 省份: {row['province_name'] or ''}",
            f"- 城市: {row['city_name'] or ''}",
            f"- 区县: {row['county_name'] or ''}",
            f"- 街道: {row['town_name'] or ''}",
            f"- 地址: {row['address'] or ''}",
            f"- 设备名称: {row['device_name'] or ''}",
            f"- 通道名称: {row['channel_name'] or ''}",
            f"- 纬度: {row['lat'] or ''}",
            f"- 经度: {row['lon'] or ''}",
            "",
            "## 事件摘要",
            f"- 自动摘要: {summary}",
            f"- 事件描述: {description}",
            f"- 告警主体: {row['alarm_body'] or ''}",
            "",
            "## 媒体引用",
            f"- 图片路径: {row['img_src_path'] or ''}",
            f"- 视频路径: {row['video_path'] or ''}",
            "",
            "## 扩展字段",
            f"```json\n{extra_json}\n```" if extra_json else "",
        ]
    ).strip() + "\n"


PROCESSING_CATEGORY_RULES = {
    "finance": ["财务", "发票", "报销", "利润", "收入", "成本", "预算", "审计", "税", "银行", "invoice", "budget"],
    "legal": ["合同", "协议", "条款", "法务", "诉讼", "仲裁", "保密", "compliance", "agreement", "contract"],
    "technical": ["api", "python", "java", "sql", "代码", "系统", "服务", "部署", "算法", "模型", "database"],
    "procurement": ["采购", "招标", "投标", "供应商", "报价", "中标", "招采", "bid", "vendor"],
    "hr": ["招聘", "简历", "绩效", "考勤", "员工", "培训", "薪酬", "attendance", "candidate"],
    "operations": ["流程", "运维", "巡检", "工单", "排班", "仓库", "物流", "告警", "incident"],
}


def _normalize_processing_text(text: str) -> str:
    cleaned = text.replace("\u3000", " ")
    cleaned = re.sub(r"[ \t]+", " ", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def _build_resource_summary(text: str, *, max_length: int = 260) -> str:
    normalized = _normalize_processing_text(text)
    if not normalized:
        return ""
    sentences = re.split(r"(?<=[。！？!?\.])\s+|\n+", normalized)
    summary_parts: list[str] = []
    total = 0
    for sentence in sentences:
        candidate = sentence.strip()
        if not candidate:
            continue
        if len(candidate) > max_length:
            candidate = candidate[: max_length - 3].rstrip() + "..."
        if total + len(candidate) > max_length and summary_parts:
            break
        summary_parts.append(candidate)
        total += len(candidate)
        if total >= max_length:
            break
    summary = " ".join(summary_parts).strip()
    return summary[:max_length]


def _classify_resource(file_name: str, file_category: Optional[str], text: str) -> str:
    lowered_text = text.lower()
    best_label = file_category or "general"
    best_score = 0
    for label, keywords in PROCESSING_CATEGORY_RULES.items():
        score = sum(1 for keyword in keywords if keyword.lower() in lowered_text)
        if score > best_score:
            best_label = label
            best_score = score
    if best_score == 0 and file_category in {"table", "spreadsheet"}:
        return "structured-data"
    return best_label


def _compute_quality_score(text: str) -> float:
    normalized = _normalize_processing_text(text)
    if not normalized:
        return 0.0
    tokens = re.findall(r"[\u4e00-\u9fff]+|[a-zA-Z0-9]+", normalized.lower())
    token_count = len(tokens)
    unique_ratio = (len(set(tokens)) / token_count) if token_count else 0.0
    length_score = min(len(normalized) / 1500, 1.0)
    line_count = normalized.count("\n") + 1
    structure_score = min(line_count / 20, 1.0)
    noise_penalty = min(normalized.count("\ufffd") / max(len(normalized), 1), 0.3)
    score = (0.5 * length_score) + (0.3 * unique_ratio) + (0.2 * structure_score) - noise_penalty
    return round(max(min(score, 1.0), 0.0) * 100, 2)


def _aggregate_embedding(vectors: list[list[float]]) -> list[float]:
    valid = [np.asarray(vector, dtype=np.float32) for vector in vectors if vector]
    if not valid:
        return []
    matrix = np.vstack(valid)
    mean_vec = matrix.mean(axis=0)
    norm = float(np.linalg.norm(mean_vec))
    if norm <= 0:
        return mean_vec.tolist()
    return (mean_vec / norm).astype(np.float32).tolist()


def _kmeans_assignments(embeddings: list[list[float]], cluster_count: int) -> list[int]:
    if not embeddings:
        return []
    vectors = np.asarray(embeddings, dtype=np.float32)
    item_count = vectors.shape[0]
    if item_count <= 1:
        return [0]
    k = max(1, min(cluster_count, item_count))
    centroids = vectors[:k].copy()
    assignments = np.zeros(item_count, dtype=np.int32)

    for _ in range(12):
        similarity = vectors @ centroids.T
        next_assignments = np.argmax(similarity, axis=1)
        if np.array_equal(assignments, next_assignments):
            break
        assignments = next_assignments
        for idx in range(k):
            members = vectors[assignments == idx]
            if members.size == 0:
                centroids[idx] = vectors[idx % item_count]
                continue
            centroid = members.mean(axis=0)
            norm = float(np.linalg.norm(centroid))
            if norm > 0:
                centroid = centroid / norm
            centroids[idx] = centroid

    return assignments.tolist()


class WorkbenchService:
    def __init__(self, db: AsyncSession):
        self.db = db

    async def import_tower_warning_sample_dataset(
        self,
        *,
        workspace_id: int,
        name: str,
        description: Optional[str] = None,
        max_records: int = 200,
    ) -> WorkbenchDataset:
        max_records = max(1, min(int(max_records), 1000))
        if not TOWER_METADATA_DB.exists():
            raise FileNotFoundError(str(TOWER_METADATA_DB))

        generated_dir = await asyncio.to_thread(
            self._build_tower_warning_sample_files,
            max_records,
        )
        return await self.create_dataset(
            workspace_id=workspace_id,
            name=name,
            description=description,
            file_paths=[str(generated_dir)],
        )

    async def create_dataset(
        self,
        *,
        workspace_id: int,
        name: str,
        description: Optional[str] = None,
        uploads: Optional[list[UploadFile]] = None,
        file_paths: Optional[list[str]] = None,
        sftp_config: Optional[dict] = None,
    ) -> WorkbenchDataset:
        uploads = uploads or []
        file_paths = [item.strip() for item in (file_paths or []) if item and item.strip()]

        if not uploads and not file_paths and not sftp_config:
            raise ValueError("Please provide files, local paths, or an SFTP source.")

        dataset = WorkbenchDataset(
            workspace_id=workspace_id,
            name=name,
            description=description,
            processing_status=ProcessingStatus.PENDING,
            source_summary={},
        )
        self.db.add(dataset)
        await self.db.flush()

        dataset_key = f"dataset_{dataset.id}"
        prepared: list[PreparedWorkbenchFile] = []
        prepared.extend(await self._prepare_uploads(uploads, dataset_key))
        prepared.extend(await self._prepare_local_paths(file_paths, dataset_key))
        if sftp_config:
            prepared.extend(await self._prepare_sftp_files(sftp_config, dataset_key))

        prepared = self._expand_archives(prepared, dataset_key)
        prepared = self._dedupe_prepared_files(prepared)
        if not prepared:
            raise ValueError("No supported files were discovered from the provided inputs.")

        await self._append_prepared_files(dataset, prepared)

        await self.db.commit()
        await self.db.refresh(dataset)
        return dataset

    def _build_tower_warning_sample_files(self, max_records: int) -> Path:
        ensure_workbench_dirs()
        target_root = WORKBENCH_IMPORT_ROOT / "tower_warning_events"
        target_root.mkdir(parents=True, exist_ok=True)

        conn = sqlite3.connect(TOWER_METADATA_DB)
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute(
            """
            SELECT
                e.event_id,
                e.asset_id,
                e.event_type,
                e.alarm_level,
                e.alarm_source,
                e.alarm_time,
                e.lat,
                e.lon,
                e.region,
                e.extra_json,
                e.summary,
                e.description,
                e.address,
                e.device_name,
                e.confidence_level,
                e.province_name,
                e.city_name,
                e.county_name,
                e.town_name,
                e.channel_name,
                e.alarm_body,
                e.algorithm_name,
                json_extract(e.extra_json, '$.app_scenarios') AS app_scenarios,
                e.video_path,
                e.img_src_path,
                a.file_path AS asset_file_path,
                a.file_name AS asset_file_name,
                a.media_type AS media_type
            FROM events e
            LEFT JOIN assets a ON a.asset_id = e.asset_id
            ORDER BY e.alarm_time DESC, e.event_id DESC
            LIMIT ?
            """,
            (max_records,),
        )
        rows = cur.fetchall()
        conn.close()

        if not rows:
            raise ValueError("No event data found in metadata.db")

        for row in rows:
            file_name = f"{row['alarm_time'] or 'event'}_{row['event_id']}.md"
            safe_name = _make_safe_name(file_name)
            output_path = target_root / safe_name
            output_path.write_text(_format_tower_event_markdown(row), encoding="utf-8")

        logger.info("Generated %s tower warning event files in %s", len(rows), target_root)
        return target_root

    async def _append_prepared_files(
        self,
        dataset: WorkbenchDataset,
        prepared: list[PreparedWorkbenchFile],
    ) -> int:
        if not prepared:
            return 0

        existing_keys = await self._get_dataset_resource_dedupe_keys(dataset.id)
        fresh = self._dedupe_prepared_files(prepared, existing_keys)
        if not fresh:
            return 0

        for item in fresh:
            self.db.add(
                WorkbenchResource(
                    dataset_id=dataset.id,
                    source_kind=item.source_kind,
                    source_uri=item.source_uri,
                    stored_path=item.stored_path,
                    file_name=item.file_name,
                    file_extension=item.file_extension,
                    file_category=item.file_category,
                    mime_type=item.mime_type,
                    file_size=item.file_size,
                    checksum=item.checksum,
                    status=ProcessingStatus.PENDING,
                    extraction_metadata=item.extraction_metadata or None,
                )
            )

        await self.db.flush()
        dataset.resource_count += len(fresh)
        dataset.source_summary = await self._build_dataset_source_summary(dataset.id)
        dataset.processing_summary = None
        dataset.processing_updated_at = None
        dataset.processing_status = ProcessingStatus.PENDING
        return len(fresh)

    async def get_dataset(self, dataset_id: int) -> Optional[WorkbenchDataset]:
        result = await self.db.execute(select(WorkbenchDataset).where(WorkbenchDataset.id == dataset_id))
        return result.scalar_one_or_none()

    async def get_dataset_detail(self, dataset_id: int) -> Optional[tuple[WorkbenchDataset, list[WorkbenchResource]]]:
        dataset = await self.get_dataset(dataset_id)
        if not dataset:
            return None
        resources_result = await self.db.execute(
            select(WorkbenchResource)
            .where(WorkbenchResource.dataset_id == dataset_id)
            .order_by(WorkbenchResource.created_at.desc(), WorkbenchResource.id.desc())
        )
        resources = list(resources_result.scalars().all())
        return dataset, resources

    async def list_datasets(self, workspace_id: int) -> list[WorkbenchDataset]:
        result = await self.db.execute(
            select(WorkbenchDataset)
            .where(WorkbenchDataset.workspace_id == workspace_id)
            .order_by(WorkbenchDataset.updated_at.desc(), WorkbenchDataset.id.desc())
        )
        return list(result.scalars().all())

    async def search_dataset(self, dataset_id: int, query: str, top_k: int = 8) -> list[WorkbenchSearchResult]:
        query_embedding = media_model_client.embed_text(query)
        result = await self.db.execute(
            select(WorkbenchChunk, WorkbenchResource)
            .join(WorkbenchResource, WorkbenchResource.id == WorkbenchChunk.resource_id)
            .where(WorkbenchChunk.dataset_id == dataset_id)
        )

        scored: list[WorkbenchSearchResult] = []
        for chunk, resource in result.all():
            vector_score = _cosine_similarity(query_embedding, chunk.embedding or [])
            keyword_score = _keyword_match_score(query, chunk.content)
            score = round((0.75 * vector_score) + (0.25 * keyword_score), 6)
            if score <= 0:
                continue
            scored.append(
                WorkbenchSearchResult(
                    dataset_id=dataset_id,
                    resource_id=resource.id,
                    chunk_id=chunk.id,
                    chunk_index=chunk.chunk_index,
                    file_name=resource.file_name,
                    source_uri=resource.source_uri,
                    score=score,
                    content=chunk.content,
                    keywords=list(chunk.keywords or []),
                    chunk_metadata=chunk.chunk_metadata,
                    parser_name=resource.parser_name,
                )
            )

        scored.sort(key=lambda item: item.score, reverse=True)
        return scored[: max(min(top_k, 20), 1)]

    async def process_dataset(
        self,
        dataset_id: int,
        *,
        extract_labels: bool = True,
        cluster_count: int = 4,
        refresh_summary: bool = True,
    ) -> tuple[WorkbenchDataset, list[WorkbenchResource]]:
        detail = await self.get_dataset_detail(dataset_id)
        if not detail:
            raise ValueError("Dataset not found")
        dataset, resources = detail
        if not resources:
            raise ValueError("Dataset has no resources")

        chunk_result = await self.db.execute(
            select(WorkbenchChunk)
            .where(WorkbenchChunk.dataset_id == dataset_id)
            .order_by(WorkbenchChunk.resource_id.asc(), WorkbenchChunk.chunk_index.asc())
        )
        chunks = list(chunk_result.scalars().all())
        if not chunks:
            raise ValueError("Dataset has no vectorized chunks yet")

        chunk_map: dict[int, list[WorkbenchChunk]] = {}
        for chunk in chunks:
            chunk_map.setdefault(chunk.resource_id, []).append(chunk)

        resource_vectors: list[list[float]] = []
        resource_rows: list[WorkbenchResource] = []
        label_counts: dict[str, int] = {}
        category_counts: dict[str, int] = {}
        quality_scores: list[float] = []
        duplicate_chunk_count = max(len(chunks) - len({chunk.content_hash for chunk in chunks if chunk.content_hash}), 0)
        processed_at = datetime.utcnow()

        for resource in resources:
            resource_chunks = chunk_map.get(resource.id, [])
            text_parts = [chunk.content for chunk in resource_chunks if chunk.content]
            combined_text = "\n\n".join(text_parts)
            cleaned_text = _normalize_processing_text(combined_text)
            labels = _extract_keywords(cleaned_text, limit=10) if extract_labels else list(resource.labels or [])
            category_label = _classify_resource(resource.file_name, resource.file_category, cleaned_text)
            summary_text = _build_resource_summary(cleaned_text)
            quality_score = _compute_quality_score(cleaned_text)
            aggregate_embedding = _aggregate_embedding([chunk.embedding or [] for chunk in resource_chunks])

            resource.labels = labels
            resource.category_label = category_label
            resource.summary_text = summary_text
            resource.quality_score = quality_score
            resource.processing_metadata = {
                "cleaned_text_length": len(cleaned_text),
                "line_count": cleaned_text.count("\n") + 1 if cleaned_text else 0,
                "token_count": len(re.findall(r"[\u4e00-\u9fff]+|[a-zA-Z0-9]+", cleaned_text)),
                "embedding_dim": len(aggregate_embedding),
                "updated_at": processed_at.isoformat(),
            }

            resource_vectors.append(aggregate_embedding or media_model_client.embed_text(summary_text or resource.file_name))
            resource_rows.append(resource)
            quality_scores.append(quality_score)
            category_counts[category_label] = category_counts.get(category_label, 0) + 1
            for label in labels:
                label_counts[label] = label_counts.get(label, 0) + 1

        assignments = _kmeans_assignments(resource_vectors, max(cluster_count, 1))
        cluster_members: dict[int, list[WorkbenchResource]] = {}
        for resource, cluster_id in zip(resource_rows, assignments):
            resource.cluster_id = int(cluster_id)
            cluster_members.setdefault(int(cluster_id), []).append(resource)

        cluster_summaries: list[dict] = []
        for cluster_id, members in sorted(cluster_members.items(), key=lambda item: item[0]):
            cluster_label_counts: dict[str, int] = {}
            for member in members:
                for label in member.labels or []:
                    cluster_label_counts[label] = cluster_label_counts.get(label, 0) + 1
            if cluster_label_counts:
                ranked_labels = sorted(cluster_label_counts.items(), key=lambda item: (-item[1], item[0]))
                cluster_label = " / ".join(label for label, _count in ranked_labels[:3])
            else:
                cluster_categories = [member.category_label or "general" for member in members]
                cluster_label = cluster_categories[0]

            for member in members:
                member.cluster_label = cluster_label

            cluster_summaries.append(
                {
                    "cluster_id": cluster_id,
                    "label": cluster_label,
                    "resource_count": len(members),
                    "resource_ids": [member.id for member in members],
                }
            )

        sorted_labels = sorted(label_counts.items(), key=lambda item: (-item[1], item[0]))
        avg_quality = round(sum(quality_scores) / len(quality_scores), 2) if quality_scores else 0.0
        summary = dataset.processing_summary if isinstance(dataset.processing_summary, dict) else {}
        if refresh_summary:
            recommendations: list[str] = []
            if avg_quality < 55:
                recommendations.append("文本质量偏低，建议补充更干净的数据源或 OCR。")
            if duplicate_chunk_count > 0:
                recommendations.append("检测到重复文本块，建议进一步做去重与版本治理。")
            if len(category_counts) <= 1 and len(resources) > 3:
                recommendations.append("当前分类较单一，建议补充标签规则或引入更细粒度业务分类。")

            summary = {
                "processed_at": processed_at.isoformat(),
                "resource_count": len(resources),
                "avg_quality_score": avg_quality,
                "duplicate_chunk_count": duplicate_chunk_count,
                "top_labels": [{"label": label, "count": count} for label, count in sorted_labels[:12]],
                "category_distribution": category_counts,
                "clusters": cluster_summaries,
                "recommendations": recommendations,
            }

        dataset.processing_summary = summary
        dataset.processing_updated_at = processed_at
        await self.db.commit()
        return await self.get_dataset_detail(dataset_id)

    async def export_dataset_payload(self, dataset_id: int) -> dict:
        detail = await self.get_dataset_detail(dataset_id)
        if not detail:
            raise ValueError("Dataset not found")
        dataset, resources = detail

        chunk_result = await self.db.execute(
            select(WorkbenchChunk)
            .where(WorkbenchChunk.dataset_id == dataset_id)
            .order_by(WorkbenchChunk.resource_id.asc(), WorkbenchChunk.chunk_index.asc())
        )
        chunks = list(chunk_result.scalars().all())

        resource_map: dict[int, list[WorkbenchChunk]] = {}
        for chunk in chunks:
            resource_map.setdefault(chunk.resource_id, []).append(chunk)

        return {
            "dataset": {
                "id": dataset.id,
                "workspace_id": dataset.workspace_id,
                "name": dataset.name,
                "description": dataset.description,
                "processing_status": dataset.processing_status.value if hasattr(dataset.processing_status, "value") else str(dataset.processing_status),
                "resource_count": dataset.resource_count,
                "processed_count": dataset.processed_count,
                "failed_count": dataset.failed_count,
                "chunk_count": dataset.chunk_count,
                "vector_count": dataset.vector_count,
                "total_text_length": dataset.total_text_length,
                "source_summary": dataset.source_summary,
                "processing_summary": dataset.processing_summary,
                "error_message": dataset.error_message,
                "last_ingested_at": dataset.last_ingested_at.isoformat() if dataset.last_ingested_at else None,
                "processing_updated_at": dataset.processing_updated_at.isoformat() if dataset.processing_updated_at else None,
                "created_at": dataset.created_at.isoformat() if dataset.created_at else None,
                "updated_at": dataset.updated_at.isoformat() if dataset.updated_at else None,
            },
            "resources": [
                {
                    "id": resource.id,
                    "source_kind": resource.source_kind.value if hasattr(resource.source_kind, "value") else str(resource.source_kind),
                    "source_uri": resource.source_uri,
                    "file_name": resource.file_name,
                    "file_extension": resource.file_extension,
                    "file_category": resource.file_category,
                    "mime_type": resource.mime_type,
                    "file_size": resource.file_size,
                    "checksum": resource.checksum,
                    "parser_name": resource.parser_name,
                    "status": resource.status.value if hasattr(resource.status, "value") else str(resource.status),
                    "chunk_count": resource.chunk_count,
                    "text_length": resource.text_length,
                    "extraction_metadata": resource.extraction_metadata,
                    "labels": list(resource.labels or []),
                    "category_label": resource.category_label,
                    "cluster_id": resource.cluster_id,
                    "cluster_label": resource.cluster_label,
                    "summary_text": resource.summary_text,
                    "quality_score": resource.quality_score,
                    "processing_metadata": resource.processing_metadata,
                    "error_message": resource.error_message,
                    "chunks": [
                        {
                            "chunk_index": chunk.chunk_index,
                            "content": chunk.content,
                            "keywords": chunk.keywords or [],
                            "chunk_metadata": chunk.chunk_metadata or {},
                        }
                        for chunk in resource_map.get(resource.id, [])
                    ],
                }
                for resource in resources
            ],
        }

    async def delete_dataset(self, dataset_id: int) -> None:
        detail = await self.get_dataset_detail(dataset_id)
        if not detail:
            return
        dataset, resources = detail

        stored_paths = {resource.stored_path for resource in resources if resource.stored_path}
        await self.db.execute(delete(WorkbenchChunk).where(WorkbenchChunk.dataset_id == dataset_id))
        await self.db.execute(delete(WorkbenchResource).where(WorkbenchResource.dataset_id == dataset_id))
        await self.db.execute(delete(WorkbenchDataset).where(WorkbenchDataset.id == dataset_id))
        await self.db.commit()

        for stored_path in stored_paths:
            self._safe_remove_file(stored_path)

        dataset_dir = WORKBENCH_EXTRACT_ROOT / f"dataset_{dataset_id}"
        if dataset_dir.exists():
            shutil.rmtree(dataset_dir, ignore_errors=True)

    async def list_subscriptions(
        self,
        workspace_id: int,
        dataset_id: Optional[int] = None,
    ) -> list[WorkbenchSubscription]:
        query = select(WorkbenchSubscription).where(WorkbenchSubscription.workspace_id == workspace_id)
        if dataset_id is not None:
            query = query.where(WorkbenchSubscription.dataset_id == dataset_id)
        query = query.order_by(WorkbenchSubscription.updated_at.desc(), WorkbenchSubscription.id.desc())
        result = await self.db.execute(query)
        return list(result.scalars().all())

    async def get_subscription(self, subscription_id: int) -> Optional[WorkbenchSubscription]:
        result = await self.db.execute(
            select(WorkbenchSubscription).where(WorkbenchSubscription.id == subscription_id)
        )
        return result.scalar_one_or_none()

    async def create_subscription(
        self,
        *,
        workspace_id: int,
        dataset_id: int,
        name: str,
        source_kind: str,
        interval_minutes: int,
        is_enabled: bool,
        local_paths: Optional[list[str]] = None,
        sftp_config: Optional[dict] = None,
    ) -> WorkbenchSubscription:
        dataset = await self.get_dataset(dataset_id)
        if not dataset or dataset.workspace_id != workspace_id:
            raise ValueError("Dataset not found in the workspace.")

        normalized_kind = str(source_kind).strip().lower()
        if normalized_kind not in {WorkbenchSourceKind.PATH.value, WorkbenchSourceKind.SFTP.value}:
            raise ValueError("Subscription source_kind must be 'path' or 'sftp'.")
        if interval_minutes <= 0:
            raise ValueError("interval_minutes must be greater than zero.")

        local_paths = [item.strip() for item in (local_paths or []) if item and item.strip()]
        if normalized_kind == WorkbenchSourceKind.PATH.value and not local_paths:
            raise ValueError("At least one local path is required.")
        if normalized_kind == WorkbenchSourceKind.SFTP.value and not sftp_config:
            raise ValueError("SFTP config is required.")

        subscription = WorkbenchSubscription(
            workspace_id=workspace_id,
            dataset_id=dataset_id,
            name=name.strip() or f"Dataset {dataset_id} subscription",
            source_kind=WorkbenchSourceKind(normalized_kind),
            is_enabled=bool(is_enabled),
            interval_minutes=int(interval_minutes),
            local_paths=local_paths or None,
            source_config=sftp_config or None,
            last_status=WorkbenchSubscriptionStatus.IDLE,
            next_run_at=datetime.utcnow() + timedelta(minutes=int(interval_minutes)) if is_enabled else None,
        )
        self.db.add(subscription)
        await self.db.commit()
        await self.db.refresh(subscription)
        return subscription

    async def update_subscription(
        self,
        subscription_id: int,
        *,
        name: Optional[str] = None,
        interval_minutes: Optional[int] = None,
        is_enabled: Optional[bool] = None,
        local_paths: Optional[list[str]] = None,
        sftp_config: Optional[dict] = None,
    ) -> WorkbenchSubscription:
        subscription = await self.get_subscription(subscription_id)
        if not subscription:
            raise ValueError("Subscription not found.")

        if name is not None:
            subscription.name = name.strip() or subscription.name
        if interval_minutes is not None:
            if interval_minutes <= 0:
                raise ValueError("interval_minutes must be greater than zero.")
            subscription.interval_minutes = int(interval_minutes)
        if is_enabled is not None:
            subscription.is_enabled = bool(is_enabled)
            if subscription.is_enabled:
                subscription.next_run_at = datetime.utcnow() + timedelta(minutes=subscription.interval_minutes)
        if local_paths is not None:
            subscription.local_paths = [item.strip() for item in local_paths if item and item.strip()] or None
        if sftp_config is not None:
            merged_config = dict(subscription.source_config or {})
            incoming_config = dict(sftp_config or {})
            if incoming_config.get("password") == "__configured__":
                incoming_config["password"] = merged_config.get("password")
            merged_config.update(incoming_config)
            subscription.source_config = merged_config or None

        if not subscription.is_enabled:
            subscription.next_run_at = None

        await self.db.commit()
        await self.db.refresh(subscription)
        return subscription

    async def delete_subscription(self, subscription_id: int) -> None:
        await self.db.execute(delete(WorkbenchSubscription).where(WorkbenchSubscription.id == subscription_id))
        await self.db.commit()

    async def run_subscription(self, subscription_id: int) -> tuple[WorkbenchSubscription, int]:
        subscription = await self.get_subscription(subscription_id)
        if not subscription:
            raise ValueError("Subscription not found.")

        dataset = await self.get_dataset(subscription.dataset_id)
        if not dataset:
            raise ValueError("Dataset not found.")

        subscription.last_status = WorkbenchSubscriptionStatus.RUNNING
        subscription.last_message = "Running"
        subscription.last_run_at = datetime.utcnow()
        await self.db.commit()

        prepared: list[PreparedWorkbenchFile] = []
        dataset_key = f"dataset_{dataset.id}"
        try:
            if subscription.source_kind == WorkbenchSourceKind.PATH:
                prepared.extend(await self._prepare_local_paths(subscription.local_paths or [], dataset_key))
            elif subscription.source_kind == WorkbenchSourceKind.SFTP:
                prepared.extend(await self._prepare_sftp_files(subscription.source_config or {}, dataset_key))
            else:
                raise ValueError("Unsupported subscription source.")

            prepared = self._expand_archives(prepared, dataset_key)
            new_count = await self._append_prepared_files(dataset, prepared)
            dataset.source_summary = await self._build_dataset_source_summary(dataset.id)
            subscription.last_status = WorkbenchSubscriptionStatus.SUCCESS
            subscription.last_message = f"Added {new_count} new resources" if new_count else "No new resources"
            subscription.last_new_resources = new_count
            subscription.next_run_at = datetime.utcnow() + timedelta(minutes=subscription.interval_minutes) if subscription.is_enabled else None
            await self.db.commit()
            await self.db.refresh(subscription)
            return subscription, new_count
        except Exception as exc:
            subscription.last_status = WorkbenchSubscriptionStatus.ERROR
            subscription.last_message = str(exc)
            subscription.last_new_resources = 0
            subscription.next_run_at = datetime.utcnow() + timedelta(minutes=subscription.interval_minutes) if subscription.is_enabled else None
            await self.db.commit()
            await self.db.refresh(subscription)
            raise

    def _summarize_sources(self, files: list[PreparedWorkbenchFile]) -> dict:
        by_source: dict[str, int] = {}
        by_category: dict[str, int] = {}
        for item in files:
            by_source[item.source_kind.value] = by_source.get(item.source_kind.value, 0) + 1
            by_category[item.file_category] = by_category.get(item.file_category, 0) + 1
        return {"source_kinds": by_source, "file_categories": by_category}

    async def _prepare_uploads(self, uploads: list[UploadFile], dataset_key: str) -> list[PreparedWorkbenchFile]:
        prepared: list[PreparedWorkbenchFile] = []
        for upload in uploads:
            if not upload.filename:
                continue
            prepared.append(await _save_upload_file(upload, WORKBENCH_UPLOAD_ROOT / dataset_key))
        return prepared

    async def _prepare_local_paths(self, file_paths: list[str], dataset_key: str) -> list[PreparedWorkbenchFile]:
        if not file_paths:
            return []
        discovered = await asyncio.to_thread(_iter_local_inputs, file_paths)
        prepared: list[PreparedWorkbenchFile] = []
        for item in discovered:
            prepared.append(await asyncio.to_thread(_copy_file_to_root, item, WORKBENCH_PATH_ROOT / dataset_key))
        return prepared

    async def _prepare_sftp_files(self, sftp_config: dict, dataset_key: str) -> list[PreparedWorkbenchFile]:
        return await asyncio.to_thread(self._download_sftp_files, sftp_config, dataset_key)

    def _download_sftp_files(self, sftp_config: dict, dataset_key: str) -> list[PreparedWorkbenchFile]:
        if paramiko is None:
            raise RuntimeError("paramiko is not installed")

        host = str(sftp_config.get("host") or "").strip()
        username = str(sftp_config.get("username") or "").strip()
        password = str(sftp_config.get("password") or "")
        port = int(sftp_config.get("port") or 22)
        recursive = bool(sftp_config.get("recursive", True))
        remote_paths = [str(item).strip() for item in (sftp_config.get("remote_paths") or []) if str(item).strip()]
        if not host or not username or not remote_paths:
            raise ValueError("SFTP host, username, and remote paths are required.")

        client = paramiko.SSHClient()
        client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        client.connect(
            hostname=host,
            port=port,
            username=username,
            password=password,
            timeout=int(settings.WORKBENCH_SFTP_TIMEOUT_SEC),
            banner_timeout=int(settings.WORKBENCH_SFTP_TIMEOUT_SEC),
            auth_timeout=int(settings.WORKBENCH_SFTP_TIMEOUT_SEC),
        )

        prepared: list[PreparedWorkbenchFile] = []
        try:
            sftp = client.open_sftp()
            remote_files = self._collect_sftp_files(sftp, remote_paths, recursive)
            target_root = WORKBENCH_SFTP_ROOT / dataset_key
            target_root.mkdir(parents=True, exist_ok=True)
            for remote_path in remote_files:
                file_name = PurePosixPath(remote_path).name
                category = _guess_file_category(file_name)
                if not category:
                    continue
                safe_name = f"{hashlib.md5(remote_path.encode('utf-8')).hexdigest()[:8]}_{_make_safe_name(file_name)}"
                target_path = target_root / safe_name
                sftp.get(remote_path, str(target_path))
                prepared.append(
                    PreparedWorkbenchFile(
                        source_kind=WorkbenchSourceKind.SFTP,
                        source_uri=f"sftp://{host}:{port}{remote_path}",
                        stored_path=str(target_path.resolve()),
                        file_name=file_name,
                        file_extension=Path(file_name).suffix.lower(),
                        file_category=category,
                        mime_type=mimetypes.guess_type(file_name)[0],
                        file_size=int(target_path.stat().st_size),
                        checksum=_compute_checksum(str(target_path)),
                        extraction_metadata={"remote_path": remote_path, "host": host, "port": port},
                    )
                )
        finally:
            client.close()

        return prepared

    def _collect_sftp_files(self, sftp, remote_paths: list[str], recursive: bool) -> list[str]:
        discovered: list[str] = []
        seen: set[str] = set()

        def walk(path: str) -> None:
            stat_result = sftp.stat(path)
            if stat.S_ISDIR(stat_result.st_mode):
                if not recursive:
                    return
                for entry in sftp.listdir_attr(path):
                    child = f"{path.rstrip('/')}/{entry.filename}"
                    if stat.S_ISDIR(entry.st_mode):
                        walk(child)
                    else:
                        if PurePosixPath(entry.filename).suffix.lower() in SUPPORTED_EXTENSIONS and child not in seen:
                            seen.add(child)
                            discovered.append(child)
                return

            if PurePosixPath(path).suffix.lower() in SUPPORTED_EXTENSIONS and path not in seen:
                seen.add(path)
                discovered.append(path)

        for remote_path in remote_paths:
            walk(remote_path)
        return discovered

    def _expand_archives(self, prepared: list[PreparedWorkbenchFile], dataset_key: str) -> list[PreparedWorkbenchFile]:
        expanded: list[PreparedWorkbenchFile] = []
        for item in prepared:
            if item.file_extension in ARCHIVE_EXTENSIONS:
                expanded.extend(_extract_archive(item, dataset_key))
                self._safe_remove_file(item.stored_path)
            else:
                expanded.append(item)
        return expanded

    def _dedupe_prepared_files(
        self,
        prepared: list[PreparedWorkbenchFile],
        existing_keys: Optional[set[str]] = None,
    ) -> list[PreparedWorkbenchFile]:
        deduped: list[PreparedWorkbenchFile] = []
        seen: set[str] = set(existing_keys or set())
        for item in prepared:
            dedupe_key = f"{item.checksum}:{item.file_name}"
            if dedupe_key in seen:
                self._safe_remove_file(item.stored_path)
                continue
            seen.add(dedupe_key)
            deduped.append(item)
        return deduped

    async def _get_dataset_resource_dedupe_keys(self, dataset_id: int) -> set[str]:
        result = await self.db.execute(
            select(WorkbenchResource.checksum, WorkbenchResource.file_name).where(
                WorkbenchResource.dataset_id == dataset_id
            )
        )
        return {
            f"{checksum}:{file_name}"
            for checksum, file_name in result.all()
            if checksum and file_name
        }

    async def _build_dataset_source_summary(self, dataset_id: int) -> dict:
        result = await self.db.execute(
            select(WorkbenchResource.source_kind, WorkbenchResource.file_category).where(
                WorkbenchResource.dataset_id == dataset_id
            )
        )
        by_source: dict[str, int] = {}
        by_category: dict[str, int] = {}
        for source_kind, file_category in result.all():
            source_value = source_kind.value if hasattr(source_kind, "value") else str(source_kind)
            if source_value:
                by_source[source_value] = by_source.get(source_value, 0) + 1
            if file_category:
                by_category[file_category] = by_category.get(file_category, 0) + 1

        return {"source_kinds": by_source, "file_categories": by_category}

    def _safe_remove_file(self, path: str) -> None:
        try:
            resolved = Path(path).resolve()
            if WORKBENCH_ROOT not in resolved.parents and resolved != WORKBENCH_ROOT:
                return
            if resolved.exists() and resolved.is_file():
                resolved.unlink()
        except Exception:
            logger.warning("Failed to remove workbench file: %s", path)


async def process_workbench_dataset_resources(db: AsyncSession, dataset_id: int) -> None:
    result = await db.execute(select(WorkbenchDataset).where(WorkbenchDataset.id == dataset_id))
    dataset = result.scalar_one_or_none()
    if not dataset:
        return

    resources_result = await db.execute(
        select(WorkbenchResource)
        .where(WorkbenchResource.dataset_id == dataset_id)
        .order_by(WorkbenchResource.id.asc())
    )
    resources = list(resources_result.scalars().all())
    if not resources:
        await reconcile_workbench_dataset_status(db, dataset_id)
        return

    dataset.processing_status = ProcessingStatus.PROCESSING
    dataset.error_message = None
    await db.commit()

    for resource in resources:
        if resource.status == ProcessingStatus.READY:
            continue
        try:
            await process_workbench_resource(db, resource)
        except Exception as exc:
            logger.error("Workbench resource processing failed: dataset=%s resource=%s error=%s", dataset_id, resource.id, exc, exc_info=True)
            resource.status = ProcessingStatus.FAILED
            resource.error_message = str(exc)
            resource.last_processed_at = datetime.utcnow()
            await db.commit()
        finally:
            await reconcile_workbench_dataset_status(db, dataset_id)


async def process_workbench_resource(db: AsyncSession, resource: WorkbenchResource) -> None:
    resource.status = ProcessingStatus.PROCESSING
    resource.error_message = None
    await db.commit()

    extracted = await asyncio.to_thread(extract_resource_chunks, resource.stored_path)
    if not extracted.chunks:
        raise RuntimeError("No text content was extracted from the file.")

    await db.execute(delete(WorkbenchChunk).where(WorkbenchChunk.resource_id == resource.id))
    chunk_rows: list[WorkbenchChunk] = []
    for index, chunk_text in enumerate(extracted.chunks):
        embedding = media_model_client.embed_text(chunk_text)
        chunk_rows.append(
            WorkbenchChunk(
                dataset_id=resource.dataset_id,
                resource_id=resource.id,
                chunk_index=index,
                content=chunk_text,
                content_hash=hashlib.sha256(chunk_text.encode("utf-8")).hexdigest(),
                embedding=embedding,
                keywords=_extract_keywords(chunk_text),
                chunk_metadata={
                    "parser": extracted.parser_name,
                    "file_category": extracted.file_category,
                    **(extracted.metadata or {}),
                },
            )
        )

    db.add_all(chunk_rows)
    resource.parser_name = extracted.parser_name
    resource.file_category = extracted.file_category
    resource.chunk_count = len(chunk_rows)
    resource.text_length = extracted.text_length
    resource.extraction_metadata = extracted.metadata or None
    resource.status = ProcessingStatus.READY
    resource.error_message = None
    resource.last_processed_at = datetime.utcnow()
    await db.commit()


async def reconcile_workbench_dataset_status(db: AsyncSession, dataset_id: int) -> None:
    result = await db.execute(select(WorkbenchDataset).where(WorkbenchDataset.id == dataset_id))
    dataset = result.scalar_one_or_none()
    if not dataset:
        return

    resource_stats = await db.execute(
        select(
            func.count(WorkbenchResource.id),
            func.sum(case((WorkbenchResource.status == ProcessingStatus.READY, 1), else_=0)),
            func.sum(case((WorkbenchResource.status == ProcessingStatus.FAILED, 1), else_=0)),
            func.sum(case((WorkbenchResource.status == ProcessingStatus.PROCESSING, 1), else_=0)),
            func.sum(case((WorkbenchResource.status == ProcessingStatus.PENDING, 1), else_=0)),
            func.sum(WorkbenchResource.text_length),
        ).where(WorkbenchResource.dataset_id == dataset_id)
    )
    total, ready_count, failed_count, processing_count, pending_count, total_text_length = resource_stats.one()

    chunk_stats = await db.execute(
        select(func.count(WorkbenchChunk.id)).where(WorkbenchChunk.dataset_id == dataset_id)
    )
    chunk_count = int(chunk_stats.scalar() or 0)

    dataset.resource_count = int(total or 0)
    dataset.processed_count = int(ready_count or 0)
    dataset.failed_count = int(failed_count or 0)
    dataset.chunk_count = chunk_count
    dataset.vector_count = chunk_count
    dataset.total_text_length = int(total_text_length or 0)

    if dataset.resource_count == 0:
        dataset.processing_status = ProcessingStatus.READY
        dataset.error_message = None
    elif int(processing_count or 0) > 0:
        dataset.processing_status = ProcessingStatus.PROCESSING
    elif int(pending_count or 0) > 0:
        dataset.processing_status = ProcessingStatus.PENDING
    elif dataset.failed_count > 0:
        dataset.processing_status = ProcessingStatus.FAILED
    else:
        dataset.processing_status = ProcessingStatus.READY

    resource_rows = await db.execute(
        select(WorkbenchResource)
        .where(WorkbenchResource.dataset_id == dataset_id)
        .order_by(WorkbenchResource.updated_at.desc())
    )
    resources = list(resource_rows.scalars().all())
    source_summary = {"source_kinds": {}, "file_categories": {}}
    for resource in resources:
        source_value = resource.source_kind.value if hasattr(resource.source_kind, "value") else str(resource.source_kind)
        if source_value:
            source_summary["source_kinds"][source_value] = source_summary["source_kinds"].get(source_value, 0) + 1
        if resource.file_category:
            source_summary["file_categories"][resource.file_category] = source_summary["file_categories"].get(resource.file_category, 0) + 1
    dataset.source_summary = source_summary
    dataset.error_message = next((resource.error_message for resource in resources if resource.error_message), None)
    dataset.last_ingested_at = max(
        (resource.last_processed_at for resource in resources if resource.last_processed_at),
        default=dataset.last_ingested_at,
    )
    await db.commit()
